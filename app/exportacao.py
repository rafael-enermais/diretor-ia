# -*- coding: utf-8 -*-
"""Diretor.IA - EXPORTACAO em planilha (.xlsx) das pendencias, para levar a uma reuniao ou mandar por e-mail. Gerada em memoria (nada e' gravado em disco).
Seguranca: texto que comeca com = + - @ vira texto puro (o relato vem de terceiros e o Excel poderia executar como formula)."""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import reports
import vinculos
from fuso import BR

NAVY, ORANGE, CINZA = "171C60", "F99D20", "F3F5F9"


def _seguro(v):
    """Evita injecao de formula: texto iniciado por = + - @ (ou tab/CR) ganha um apostrofo na frente."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _dh(v):
    """datetime com fuso -> datetime de Brasilia sem fuso (o Excel nao aceita fuso). date e vazio passam."""
    if isinstance(v, dt.datetime):
        if v.tzinfo is not None:
            v = v.astimezone(BR).replace(tzinfo=None)
        return v.replace(microsecond=0)
    return v


def _cabecalho(ws, titulos, larguras):
    ws.append(titulos)
    for i, (t, w) in enumerate(zip(titulos, larguras), 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24


def pendencias_xlsx(itens: list, historico: list, catalogo: list, gerado_em: dt.datetime, filtro_texto: str = "") -> bytes:
    """itens = os itens a exportar (ja filtrados na tela); historico = linhas de report_item_historico desses itens; catalogo = itens do catalogo (para o rotulo do vinculo)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pendências"
    cols = ["Nº", "Semana", "Colaborador", "Tipo", "Urgência", "Situação", "Título", "Descrição", "Nota do Diretor", "Vínculo", "Origem", "Recebido em", "Resolvido em", "Atualizado em"]
    _cabecalho(ws, cols, [7, 12, 22, 14, 11, 18, 36, 60, 40, 30, 14, 17, 17, 17])
    for i in itens:
        ws.append([i["id"], i["semana_de"], _seguro(i["colaborador"]), reports.TIPOS.get(i["tipo"], i["tipo"]), reports.URGENCIAS.get(i["urgencia"], i["urgencia"]),
                   reports.SITUACOES.get(i["situacao"], i["situacao"]), _seguro(i.get("titulo") or ""), _seguro(i.get("descricao") or ""), _seguro(i.get("nota_diretor") or ""),
                   _seguro(vinculos.rotulo(i.get("vinculo_item_id"), catalogo)), _seguro(i.get("origem") or ""), _dh(i.get("recebido_em")), _dh(i.get("resolvido_em")),
                   _dh(i.get("atualizado_em"))])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = "DD/MM/YYYY"
        for c in (12, 13, 14):
            ws.cell(row=r, column=c).number_format = "DD/MM/YYYY HH:MM"
        for c in (7, 8, 9, 10):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    rs = wb.create_sheet("Resumo")
    rs.column_dimensions["A"].width, rs.column_dimensions["B"].width = 34, 14
    rs["A1"] = "Diretor.IA — Reports e pendências"
    rs["A1"].font = Font(bold=True, size=14, color=NAVY)
    rs["A2"] = f"Gerado em {gerado_em.astimezone(BR).strftime('%d/%m/%Y %H:%M')} (horário de Brasília)"
    rs["A3"] = _seguro(f"Filtro aplicado: {filtro_texto}") if filtro_texto else "Filtro aplicado: nenhum (todos os itens)"
    k = reports.kpis(itens, gerado_em.astimezone(BR).date())
    linha = 5
    for rot, v in (("Itens na planilha", len(itens)), ("Em aberto (nova ou em acompanhamento)", k["abertas"]), ("Urgentes em aberto (alta ou crítica)", k["urgentes"]),
                   ("Críticas em aberto", k["criticas"]), ("Resolvidas nos últimos 30 dias", k["resolvidas_30d"]), ("Colaboradores com item em aberto", k["colaboradores_com_aberta"]),
                   ("Tempo médio para resolver (dias)", k["dias_medios_para_resolver"] if k["dias_medios_para_resolver"] is not None else "—")):
        rs.cell(row=linha, column=1, value=rot)
        rs.cell(row=linha, column=2, value=v).font = Font(bold=True, color=NAVY)
        linha += 1
    linha += 1
    rs.cell(row=linha, column=1, value="Em aberto por colaborador e urgência").font = Font(bold=True, color=NAVY)
    linha += 1
    ab = reports.abertas_por_colaborador(itens)
    if len(ab):
        for j, t in enumerate(["Colaborador"] + list(ab.columns), 1):
            c = rs.cell(row=linha, column=j, value=t)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
            if j > 2:
                rs.column_dimensions[get_column_letter(j)].width = 12
        for nome, lin in ab.iterrows():
            linha += 1
            rs.cell(row=linha, column=1, value=_seguro(nome))
            for j, v in enumerate(lin.tolist(), 2):
                rs.cell(row=linha, column=j, value=int(v))
    else:
        rs.cell(row=linha, column=1, value="Nenhum item em aberto.")

    hs = wb.create_sheet("Histórico")
    _cabecalho(hs, ["Nº do item", "Quando", "Quem", "O que mudou", "De", "Para"], [10, 17, 30, 30, 36, 36])
    for h in historico:
        hs.append([h["item_id"], _dh(h["em"]), _seguro(h["por"]), reports.ROTULO_CAMPO.get(h["campo"], h["campo"]), _seguro(h.get("de") or ""), _seguro(h.get("para") or "")])
    for r in range(2, hs.max_row + 1):
        hs.cell(row=r, column=2).number_format = "DD/MM/YYYY HH:MM"
    wb.properties.creator, wb.properties.title = "Diretor.IA", "Reports e pendências"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
