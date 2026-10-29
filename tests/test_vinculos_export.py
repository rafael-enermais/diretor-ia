# -*- coding: utf-8 -*-
import datetime as dt
import io

from openpyxl import load_workbook

import exportacao
import reports
import vinculos as V
from fakes import BancoFalso

CAT = [{"item_id": "I27", "tipo": "indicador", "nome": "Contas a pagar"}, {"item_id": "I22", "tipo": "indicador", "nome": "Contas a receber"},
       {"item_id": "I03", "tipo": "indicador", "nome": "Margem de contribuição"}, {"item_id": "E12", "tipo": "entrega", "nome": "Conciliação bancária diária"},
       {"item_id": "I40", "tipo": "indicador", "nome": "Receita"}]
HOJE = dt.date(2026, 10, 29)


def test_sugerir_por_codigo_nome_e_palavras():
    s = V.sugerir("Atraso grande em I27 e dificuldade na E12", CAT)
    assert [x["item_id"] for x in s] == ["E12", "I27"] and all(x["confianca"] == "alta" and "cita" in x["motivo"] for x in s)
    s = V.sugerir("Fornecedores cobrando: contas a pagar vencidas", CAT)
    assert s[0]["item_id"] == "I27" and s[0]["confianca"] == "alta" and "nome" in s[0]["motivo"]
    s = V.sugerir("Não consegui fazer a conciliação do banco todos os dias", CAT)                 # 1 palavra em comum so': nao sugere
    assert s == []
    s = V.sugerir("A conciliação bancária ficou atrasada", CAT)
    assert s[0]["item_id"] == "E12" and s[0]["confianca"] == "media" and "bancaria" in s[0]["motivo"]


def test_sugerir_nao_confunde_conjuncao_nem_inventa():
    assert V.sugerir("Preciso de e 12 pessoas e I99 que não existe", CAT) == []                  # 'e 12' minusculo nao e' codigo; I99 nao esta no catalogo
    assert V.sugerir("", CAT) == [] and V.sugerir(None, CAT) == []
    assert V.sugerir("Receita caiu", CAT)[0]["item_id"] == "I40"                                   # nome curto inteiro no texto
    assert V.sugerir("Receita caiu", CAT) and V.sugerir("tudo certo por aqui", CAT) == []


def test_sugerir_respeita_recusados_e_limite():
    s = V.sugerir("I27 contas a pagar", CAT, {"I27"})
    assert s == []
    muitos = [{"item_id": f"I{n:02d}", "tipo": "indicador", "nome": "X"} for n in range(10, 20)]
    assert len(V.sugerir(" ".join(f"I{n:02d}" for n in range(10, 20)), muitos, limite=3)) == 3
    assert V.recusados("I27,E12,lixo,I9") == {"I27", "E12"} and V.valido("E01") and not V.valido("X01") and not V.valido(None)
    assert V.rotulo("I27", CAT) == "I27 · Contas a pagar" and V.rotulo("I99", CAT) == "I99" and V.rotulo(None, CAT) == ""


def test_validar_vinculo_so_aceita_catalogo():
    ids = {c["item_id"] for c in CAT}
    d, av = reports.validar({"colaborador": "Fulana", "semana_de": dt.date(2026, 10, 26), "vinculo_item_id": "i27"}, HOJE, ids)
    assert d["vinculo_item_id"] == "I27" and not av
    d, av = reports.validar({"colaborador": "Fulana", "semana_de": dt.date(2026, 10, 26), "vinculo_item_id": "I99"}, HOJE, ids)
    assert d["vinculo_item_id"] is None and "I99" in av[0]
    d, av = reports.validar({"colaborador": "Fulana", "semana_de": dt.date(2026, 10, 26), "vinculo_item_id": "abc"}, HOJE)
    assert d["vinculo_item_id"] is None
    d, av = reports.validar({"colaborador": "Fulana", "semana_de": dt.date(2026, 10, 26), "vinculo_item_id": ""}, HOJE, ids)
    assert d["vinculo_item_id"] is None and not av


def test_atualizar_grava_vinculo_com_historico_e_recusar():
    b = BancoFalso()
    antigo = {c: None for c in reports.CAMPOS}
    antigo.update(situacao="nova", tipo="outro", urgencia="media", colaborador="Fulana", semana_de=dt.date(2026, 10, 26))
    novo = dict(antigo, vinculo_item_id="I27")
    assert reports.atualizar(b, "u@x.com", 3, antigo, novo) == 1
    tx = b.transacoes[0]
    assert "vinculo_item_id = %s" in tx[0][0] and tx[1][1][2] == "vinculo_item_id" and tx[1][1][4] == "I27"
    b2 = BancoFalso()
    reports.recusar_vinculo(b2, "u@x.com", 3, "E12", "I27")
    upd, ins = b2.transacoes[0]
    assert upd[1][0] == "E12,I27" and ins[1][2:] == ("vinculo_recusado", "I27") and "DELETE" not in (upd[0] + ins[0]).upper()
    try:
        reports.recusar_vinculo(b2, "u@x.com", 3, "", "xx")
    except ValueError:
        pass
    else:
        raise AssertionError("vinculo invalido deveria ser recusado")


def _item(i, colab="Fulana Silva", urg="media", sit="nova", **extra):
    base = {"id": i, "origem": "manual", "colaborador": colab, "semana_de": dt.date(2026, 10, 26), "tipo": "dificuldade", "urgencia": urg, "situacao": sit, "titulo": "t",
            "descricao": "d", "nota_diretor": None, "recebido_em": dt.datetime(2026, 10, 27, 12, tzinfo=dt.timezone.utc), "resolvido_em": None,
            "atualizado_em": dt.datetime(2026, 10, 27, 12, tzinfo=dt.timezone.utc), "vinculo_item_id": None}
    base.update(extra)
    return base


def test_exportar_planilha_abas_dados_e_protecao_de_formula():
    itens = [_item(1, urg="critica", titulo="=HYPERLINK(\"http://x\")", vinculo_item_id="I27"), _item(2, "Beto Souza", sit="resolvida", resolvido_em=dt.datetime(2026, 10, 28, 12, tzinfo=dt.timezone.utc),
                                                                                                 descricao="+55 11 9999")]
    hist = [{"item_id": 1, "em": dt.datetime(2026, 10, 27, 12, tzinfo=dt.timezone.utc), "por": "u@x.com", "campo": "urgencia", "de": "Média", "para": "Crítica"}]
    dados = exportacao.pendencias_xlsx(itens, hist, CAT, dt.datetime(2026, 10, 29, 15, 0, tzinfo=dt.timezone.utc), "Situação: Nova")
    wb = load_workbook(io.BytesIO(dados))
    assert wb.sheetnames == ["Pendências", "Resumo", "Histórico"]
    ws = wb["Pendências"]
    assert ws.max_row == 3 and ws["A2"].value == 1 and ws["E2"].value == "Crítica" and ws["F3"].value == "Resolvida"
    assert ws["G2"].value.startswith("'=") and ws["H3"].value.startswith("'+") and ws["J2"].value == "I27 · Contas a pagar"        # nada vira formula
    assert ws["L2"].value == dt.datetime(2026, 10, 27, 9, 0) and ws["B2"].value == dt.datetime(2026, 10, 26)                 # hora de Brasilia (UTC-3)
    rs = wb["Resumo"]
    textos = [c.value for row in rs.iter_rows() for c in row if c.value is not None]
    assert "Filtro aplicado: Situação: Nova" in textos and "Fulana Silva" in textos and any("29/10/2026 12:00" in str(t) for t in textos)
    assert wb["Histórico"]["D2"].value == "Urgência" and wb["Histórico"]["F2"].value == "Crítica"
    vazio = load_workbook(io.BytesIO(exportacao.pendencias_xlsx([], [], CAT, dt.datetime(2026, 10, 29, 15, tzinfo=dt.timezone.utc))))
    assert vazio["Pendências"].max_row == 1 and any(c.value == "Nenhum item em aberto." for row in vazio["Resumo"].iter_rows() for c in row)
